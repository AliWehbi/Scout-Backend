"""
The "Download sample" spreadsheet for the member import. Built from the organization's own
units and custom fields, with headers in the user's language that the importer recognizes.
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from apps.units.models import Unit

from .importer import organization_custom_fields
from .models import FieldType

# target: (English header, Arabic header, required, English help, Arabic help, example 1, example 2)
# Example values per language are given as (en, ar) where they differ.
COLUMNS = [
    ("first_name", "First name", "الاسم الأول (بالإنجليزية)", True, "Required.", "إلزامي.", "Elie", "Jana"),
    ("last_name", "Last name", "اسم العائلة (بالإنجليزية)", False, "", "", "Mansour", "Khalil"),
    ("first_name_ar", "First name (Arabic)", "الاسم الأول بالعربية", False, "", "", "إيلي", "جنى"),
    ("last_name_ar", "Last name (Arabic)", "اسم العائلة بالعربية", False, "", "", "منصور", "خليل"),
    (
        "date_of_birth",
        "Date of birth",
        "تاريخ الميلاد",
        False,
        "DD/MM/YYYY, e.g. 02/06/2017.",
        "اليوم/الشهر/السنة، مثلاً 02/06/2017.",
        date(2017, 6, 2),
        date(2016, 3, 14),
    ),
    ("gender", "Gender", "الجنس", False, "Male or Female.", "ذكر أو أنثى.", ("Male", "ذكر"), ("Female", "أنثى")),
    (
        "kind",
        "Type",
        "النوع",
        False,
        "Youth or Adult. Empty means Youth.",
        "شاب أو بالغ. الفارغ يعني شاب.",
        ("Youth", "شاب"),
        ("Youth", "شاب"),
    ),
    ("unit", "Unit", "الوحدة", False, "One of your units: {units}.", "إحدى وحداتكم: {units}.", "{unit}", "{unit}"),
    (
        "position",
        "Position",
        "المنصب",
        False,
        "Member, Patrol leader, Assistant or Leader. Empty means Member.",
        "عضو أو عريف أو قائد مساعد أو قائد. الفارغ يعني عضو.",
        ("Member", "عضو"),
        ("Patrol leader", "عريف"),
    ),
    ("phone", "Phone", "الهاتف", False, "", "", "", ""),
    ("email", "Email", "البريد الإلكتروني", False, "", "", "", ""),
    ("address", "Address", "العنوان", False, "", "", ("Achrafieh, Beirut", "الأشرفية، بيروت"), ""),
    (
        "guardian_name",
        "Parent name",
        "اسم ولي الأمر",
        False,
        "",
        "",
        ("Rita Mansour", "ريتا منصور"),
        ("Sami Khalil", "سامي خليل"),
    ),
    ("guardian_phone", "Parent phone", "هاتف ولي الأمر", False, "", "", "+961 3 220 118", "+961 71 404 002"),
    ("guardian_email", "Parent email", "بريد ولي الأمر", False, "", "", "rita@example.com", ""),
    (
        "guardian_relationship",
        "Relationship",
        "صلة القرابة",
        False,
        "Mother, Father, Guardian, Grandparent, Sibling or Other.",
        "أم أو أب أو ولي أمر أو جد أو جدة أو أخ أو أخت أو آخر.",
        ("Mother", "أم"),
        ("Father", "أب"),
    ),
    (
        "allergies",
        "Allergies",
        "الحساسية",
        False,
        "Separate several with commas.",
        "افصل بينها بفواصل.",
        ("Pollen", "حبوب اللقاح"),
        "",
    ),
]

LISTS = {
    "gender": (["Male", "Female"], ["ذكر", "أنثى"]),
    "kind": (["Youth", "Adult"], ["شاب", "بالغ"]),
    "position": (["Member", "Patrol leader", "Assistant", "Leader"], ["عضو", "عريف", "قائد مساعد", "قائد"]),
    "guardian_relationship": (
        ["Mother", "Father", "Guardian", "Grandparent", "Sibling", "Other"],
        ["أم", "أب", "ولي أمر", "جد", "جدة", "أخ", "أخت", "آخر"],
    ),
}

TEXT = {
    "en": {
        "sheet": "Members",
        "help_sheet": "Instructions",
        "column": "Column",
        "required": "Required",
        "what": "What to enter",
        "yes": "Yes",
        "intro": "Fill one member per row on the “Members” sheet. Replace or delete the two example rows. "
        "Any column can be left out; column names can be changed — you can match them during the import.",
        "yesno": "Yes or No.",
        "number": "A number.",
        "custom_date": "A date, e.g. 2026-09-14.",
        "options": "One of: {options}.",
        "multi": "One or more of: {options}, separated by commas.",
    },
    "ar": {
        "sheet": "الأعضاء",
        "help_sheet": "التعليمات",
        "column": "العمود",
        "required": "إلزامي",
        "what": "ما يجب إدخاله",
        "yes": "نعم",
        "intro": "أدخل عضواً واحداً في كل صف في ورقة «الأعضاء». استبدل الصفين المثالين أو احذفهما. "
        "يمكن حذف أي عمود أو تغيير اسمه، ويمكنك مطابقة الأعمدة أثناء الاستيراد.",
        "yesno": "نعم أو لا.",
        "number": "رقم.",
        "custom_date": "تاريخ، مثلاً 2026-09-14.",
        "options": "أحد الخيارات: {options}.",
        "multi": "خيار أو أكثر من: {options}، مفصولة بفواصل.",
    },
}


def _pick(value, ar: bool):
    return value[1 if ar else 0] if isinstance(value, tuple) else value


def build_sample(organization, lang: str) -> bytes:
    ar = lang == "ar"
    text = TEXT["ar" if ar else "en"]
    units = [(u.name_ar if ar and u.name_ar else u.name_en) for u in Unit.objects.filter(organization=organization)]
    unit_list = ", ".join(units) or "—"
    first_unit = units[0] if units else ""

    headers, required, helps, rows = [], [], [], [[], []]
    validations = []  # (column index, options)
    for target, en, ar_header, is_required, help_en, help_ar, *examples in COLUMNS:
        headers.append(ar_header if ar else en)
        required.append(is_required)
        helps.append((help_ar if ar else help_en).format(units=unit_list))
        for row, example in zip(rows, examples, strict=True):
            row.append(first_unit if example == "{unit}" else _pick(example, ar))
        if target in LISTS:
            validations.append((len(headers), LISTS[target][1 if ar else 0]))
        elif target == "unit" and units:
            validations.append((len(headers), units))

    for field in organization_custom_fields(organization):
        headers.append(field.label_ar if ar and field.label_ar else field.label_en)
        required.append(False)
        options = ", ".join(field.options or [])
        helps.append(
            {
                FieldType.YES_NO: text["yesno"],
                FieldType.NUMBER: text["number"],
                FieldType.DATE: text["custom_date"],
                FieldType.CHOICE: text["options"].format(options=options),
                FieldType.MULTI: text["multi"].format(options=options),
            }.get(field.field_type, "")
        )
        example = (field.options or [""])[0] if field.field_type in (FieldType.CHOICE, FieldType.MULTI) else ""
        for row in rows:
            row.append(example)
        if field.field_type == FieldType.CHOICE and field.options:
            validations.append((len(headers), field.options))
        elif field.field_type == FieldType.YES_NO:
            validations.append((len(headers), ["نعم", "لا"] if ar else ["Yes", "No"]))

    book = Workbook()
    sheet = book.active
    sheet.title = text["sheet"]
    sheet.sheet_view.rightToLeft = ar
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="E8EEF4")
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font = bold
        cell.fill = fill
        sheet.column_dimensions[get_column_letter(index)].width = max(14, len(str(header)) + 4)
    date_column = [c[0] for c in COLUMNS].index("date_of_birth") + 1
    for r in range(2, 502):
        sheet.cell(row=r, column=date_column).number_format = "DD/MM/YYYY"
    sheet.freeze_panes = "A2"
    for index, options in validations:
        if any("," in o or '"' in o for o in options):
            continue
        formula = ",".join(options)
        if len(formula) > 250:  # Excel's limit for inline lists
            continue
        letter = get_column_letter(index)
        rule = DataValidation(type="list", formula1=f'"{formula}"', allow_blank=True, showErrorMessage=False)
        rule.add(f"{letter}2:{letter}501")
        sheet.add_data_validation(rule)

    guide = book.create_sheet(text["help_sheet"])
    guide.sheet_view.rightToLeft = ar
    guide.append([text["intro"]])
    guide.merge_cells("A1:C1")
    guide["A1"].alignment = Alignment(wrap_text=True, vertical="top")
    guide.row_dimensions[1].height = 48
    guide.append([])
    guide.append([text["column"], text["required"], text["what"]])
    for cell in guide[3]:
        cell.font = bold
        cell.fill = fill
    for header, is_required, help_text in zip(headers, required, helps, strict=True):
        guide.append([header, text["yes"] if is_required else "", help_text])
    guide.column_dimensions["A"].width = 30
    guide.column_dimensions["B"].width = 12
    guide.column_dimensions["C"].width = 80

    out = io.BytesIO()
    book.save(out)
    return out.getvalue()
