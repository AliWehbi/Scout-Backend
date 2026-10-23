import io
import json
import zipfile
from io import BytesIO

import pytest
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from PIL import Image

from apps.customfields.models import CustomField
from apps.members.models import Member
from apps.units.models import UnitMembership

EXPORT = "/api/v1/members/export/"
TEMPLATES = "/api/v1/document-templates/"


@pytest.fixture(autouse=True)
def _private_root(settings, tmp_path):
    settings.PRIVATE_MEDIA_ROOT = tmp_path / "private"
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def setup(world, client_for):
    org = world.org("1st Cedar Scout Group")
    cubs = world.unit(org, "Cubs")
    admin = world.member(org, world.user("rania"), owner=True)
    for first in ("Jad", "=1+1"):  # a name that looks like a formula must stay text
        m = Member.objects.create(organization=org, first_name=first)
        UnitMembership.objects.create(member=m, unit=cubs)
    CustomField.objects.create(
        organization=org, entity="organization", key="phone", code="phone", label_en="Phone",
        default={"en": "+961 1 234 567", "ar": "٠١ ٢٣٤ ٥٦٧"},
    )  # fmt: skip
    return dict(org=org, admin=admin, rania=client_for(admin.user, org))


def workbook(cells: dict, *, styles=None) -> bytes:
    wb = Workbook()
    ws = wb.active
    for ref, value in cells.items():
        ws[ref] = value
    for ref, color in (styles or {}).items():
        ws[ref].fill = PatternFill("solid", fgColor=color)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def upload(client, data: bytes, filename="list.xlsx", name="Our list"):
    file = SimpleUploadedFile(filename, data, content_type="application/octet-stream")
    return client.post(TEMPLATES, {"kind": "members_list", "name": name, "file": file}, format="multipart")


def error(res) -> str:
    return " ".join(res.data["errors"]["file"])


def export(client, **params):
    base = {"fields": "full_name", "groups": "leaders,members", "statuses": "active", "lang": "en"}
    res = client.get(EXPORT, {**base, **params})
    assert res.status_code == 200, getattr(res, "data", res)
    return load_workbook(BytesIO(res.content)).active


def test_upload_checks(setup):
    rania, code = setup["rania"], setup["org"].code
    ok = workbook({"A1": "{{ org.name }}", "A3": "{{ table }}"})
    assert upload(rania, ok).status_code == 201
    assert "Excel" in error(upload(rania, ok, filename="list.docx"))
    assert "Unknown" in error(upload(rania, workbook({"A1": "{{ org.nmae }}", "A3": "{{ table }}"})))
    assert "Unknown" in error(upload(rania, workbook({"A1": "{{ event.title }}", "A3": "{{ table }}"})))
    assert "table" in error(upload(rania, workbook({"A1": "{{ org.name }}"})))
    assert "once" in error(upload(rania, workbook({"A1": "{{ table }}", "A5": "{{ table }}"})))
    assert "alone" in error(upload(rania, workbook({"A1": "Members: {{ table }}"})))
    assert upload(rania, workbook({"A1": f"Call {{{{ org.{code}_phone }}}}", "A3": "{{table}}"})).status_code == 201

    # Macro-enabled workbooks and links to other files are refused before anything parses them.
    def rezip(data, edit):
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w") as dst:
            for info in src.infolist():
                dst.writestr(info.filename, edit(info.filename, src.read(info)))
        return out.getvalue()

    macro = rezip(
        ok,
        lambda n, b: b.replace(b"sheet.main+xml", b"sheet.macroEnabled.main+xml") if n == "[Content_Types].xml" else b,
    )
    assert "macro" in error(upload(rania, macro))
    linked = rezip(
        ok,
        lambda n, b: (
            b.replace(
                b"</Relationships>",
                b'<Relationship Id="rX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" Target="externalLinks/externalLink1.xml"/></Relationships>',
            )
            if n == "xl/_rels/workbook.xml.rels"
            else b
        ),
    )  # noqa: E501
    assert upload(rania, linked).status_code == 400


def test_export_fills_the_template(setup, client_for):
    rania, org = setup["rania"], setup["org"]
    org.logo.save("logo.png", ContentFile(_png()), save=True)
    data = workbook(
        {
            "A1": "{{ org.logo }}",
            "B1": "{{ org.name }}",
            "B2": f"Phone: {{{{ org.{org.code}_phone }}}}",
            "B3": "{{ count }} members, exported by {{ generated_by }}",
            "A5": "{{ table }}",
            "A7": "Signed: ________",
        },
        styles={"A5": "2F5D3A"},
    )
    template = upload(rania, data).data["id"]  # the first one becomes the default

    ws = export(rania)
    assert ws["B1"].value == "1st Cedar Scout Group"
    assert ws["B2"].value == "Phone: +961 1 234 567"
    assert ws["B3"].value == "2 members, exported by rania"
    assert ws["A1"].value is None and len(ws._images) == 1  # the logo
    # The header goes where {{ table }} was, in its style; the rows below; the footer moves down.
    assert ws["A5"].value == "Full name" and ws["A5"].fill.fgColor.rgb.endswith("2F5D3A")
    assert sorted([ws["A6"].value, ws["A7"].value]) == ["=1+1", "Jad"]
    assert ws["A8"].value is None and ws["A9"].value == "Signed: ________"
    formula = ws["A6"] if ws["A6"].value == "=1+1" else ws["A7"]
    assert formula.data_type == "s"  # stays text

    # Arabic: the organization field in Arabic, the column titles in Arabic.
    ws = export(rania, lang="ar")
    assert ws["B2"].value == "Phone: ٠١ ٢٣٤ ٥٦٧" and ws["A5"].value == "الاسم الكامل"

    # Values changed for this file only.
    rows = rania.get(f"{EXPORT}values/", {"template": template}).data
    assert [(r["template_code"], r["source"]) for r in rows] == [(f"org.{org.code}_phone", "organization")]
    values = json.dumps({f"org.{org.code}_phone": "+961 3 999 999"})
    assert export(rania, values=values)["B2"].value == "Phone: +961 3 999 999"
    assert CustomField.objects.get(key="phone").default["en"] == "+961 1 234 567"
    bad = rania.get(EXPORT, {"fields": "full_name", "values": json.dumps({"event.x": "1"})})
    assert bad.status_code == 400

    # `template=` asks for the built-in layout.
    ws = export(rania, template="")
    assert ws["A1"].value == "Full name"
    assert rania.get(f"{EXPORT}values/", {"template": ""}).data == []


def _png() -> bytes:
    out = BytesIO()
    Image.new("RGB", (120, 60), "green").save(out, "PNG")
    return out.getvalue()
