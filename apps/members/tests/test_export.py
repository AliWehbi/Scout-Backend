from io import BytesIO

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.members.models import CustomField, Guardian, GuardianLink, MedicalInfo, Member
from apps.units.models import UnitMembership

URL = "/api/v1/members/export/"


@pytest.fixture
def org(world):
    o = world.org()
    cubs, scouts = world.unit(o, "Cubs"), world.unit(o, "Scouts")
    admin = world.member(o, world.user("rania"), owner=True)
    karim = world.member(o, world.user("karim"), roles=[("unit_leader", cubs)])
    nour = world.member(o, world.user("nour"), roles=[("assistant_leader", cubs)])

    def person(first, unit=None, position="member", **extra):
        m = Member.objects.create(organization=o, first_name=first, **extra)
        if unit:
            UnitMembership.objects.create(member=m, unit=unit, position=position)
        return m

    zeina = person("Zeina", cubs, "leader", kind="adult", first_name_ar="زينة")
    jad = person("Jad", cubs, phone="+961 3 111")
    omar = person("Omar", scouts)
    sami = person("Sami", scouts, "assistant", kind="adult")
    old = person("Old", cubs, status="alumni")
    guardian = Guardian.objects.create(organization=o, first_name="Maya", last_name="Saliba", phone="+961 3 222")
    GuardianLink.objects.create(member=jad, guardian=guardian, relationship="mother", is_primary=True)
    MedicalInfo.objects.create(member=jad, allergies=[{"name": "Peanuts", "severity": "severe"}], blood_type="O+")
    MedicalInfo.objects.create(member=omar, blood_type="A-")
    return dict(org=o, cubs=cubs, scouts=scouts, admin=admin, karim=karim, nour=nour,
                zeina=zeina, jad=jad, omar=omar, sami=sami, old=old)  # fmt: skip


def sheet(res):
    assert res.status_code == 200, getattr(res, "data", res)
    assert res["Content-Type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    ws = load_workbook(BytesIO(res.content)).active
    return [[c.value for c in row] for row in ws.iter_rows()], ws


def test_fields_catalog_depends_on_permissions(org, client_for):
    CustomField.objects.create(organization=org["org"], key="size", label_en="Uniform size", label_ar="المقاس")
    CustomField.objects.create(organization=org["org"], key="vaccine", label_en="Vaccinated", section="medical")
    admin = client_for(org["admin"].user, org["org"]).get(f"{URL}fields/").data
    keys = [f["key"] for f in admin["fields"]]
    assert {"group", "full_name", "units", "allergies", "custom__size", "custom__vaccine"} <= set(keys)
    assert admin["default"][0] == "group"

    nour = client_for(org["nour"].user, org["org"]).get(f"{URL}fields/").data  # no medical permission
    keys = {f["key"] for f in nour["fields"]}
    assert "allergies" not in keys and "custom__vaccine" not in keys and "custom__size" in keys


def test_leaders_first_then_members_in_chosen_order(org, client_for):
    client = client_for(org["admin"].user, org["org"])
    rows, ws = sheet(client.get(URL, {"fields": "full_name,group,units,position", "lang": "en"}))
    assert rows[0] == ["Full name", "Leader / member", "Unit", "Position"]
    assert [r[0] for r in rows[1:]] == ["Sami", "Zeina", "Jad", "Omar"]  # alumni left out by default
    assert rows[1][1:] == ["Leader", "Scouts", "Assistant leader"]
    assert rows[3][1] == "Member"
    assert ws.freeze_panes == "A2"

    rows, _ = sheet(client.get(URL, {"fields": "full_name", "groups": "leaders"}))
    assert [r[0] for r in rows[1:]] == ["Sami", "Zeina"]
    rows, _ = sheet(client.get(URL, {"fields": "full_name", "groups": "members", "statuses": "active,alumni"}))
    assert [r[0] for r in rows[1:]] == ["Jad", "Old", "Omar"]


def test_unit_selection_and_unit_columns(org, client_for):
    UnitMembership.objects.create(member=org["jad"], unit=org["scouts"], position="leader")
    client = client_for(org["admin"].user, org["org"])
    # Only Cubs: Jad is a member there, and the unit columns only show Cubs
    rows, _ = sheet(client.get(URL, {"fields": "full_name,group,units", "units": org["cubs"].pk}))
    assert rows[1:] == [["Zeina", "Leader", "Cubs"], ["Jad", "Member", "Cubs"]]
    # Both units: Jad leads Scouts, so he is a leader, with both units listed
    units = f"{org['cubs'].pk},{org['scouts'].pk}"
    rows, _ = sheet(client.get(URL, {"fields": "full_name,group,units", "units": units}))
    assert ["Jad", "Leader", "Cubs, Scouts"] in rows


def test_arabic_sheet(org, client_for):
    client = client_for(org["admin"].user, org["org"])
    rows, ws = sheet(
        client.get(URL, {"fields": "full_name_ar,group", "groups": "leaders", "units": org["cubs"].pk, "lang": "ar"})
    )
    assert ws.sheet_view.rightToLeft
    assert rows == [["الاسم الكامل (عربي)", "قائد / عضو"], ["زينة", "قائد"]]


def test_medical_columns_follow_unit_scoped_permission(org, client_for):
    karim = client_for(org["karim"].user, org["org"])  # medical for Cubs only, sees only Cubs
    rows, _ = sheet(karim.get(URL, {"fields": "full_name,allergies,blood_type"}))
    assert ["Jad", "Peanuts (severe)", "O+"] in rows
    assert all(r[0] != "Omar" for r in rows)

    nour = client_for(org["nour"].user, org["org"])
    res = nour.get(URL, {"fields": "full_name,allergies"})
    assert res.status_code == 400 and "allergies" in str(res.data)

    admin = client_for(org["admin"].user, org["org"])
    sheet(admin.get(URL, {"fields": "full_name,blood_type"}))
    log = AuditLog.objects.filter(action="person.exported").latest("created_at")
    assert log.changes["medical_fields"] == ["blood_type"] and log.changes["rows"] == 4


def test_scope_and_validation(org, world, client_for):
    karim = client_for(org["karim"].user, org["org"])
    assert karim.get(URL, {"fields": "full_name", "units": org["scouts"].pk}).status_code == 403
    assert karim.get(URL, {"fields": ""}).status_code == 400
    assert karim.get(URL, {"fields": "full_name", "groups": "everyone"}).status_code == 400
    assert karim.get(URL, {"fields": "full_name", "units": "abc"}).status_code == 400
    other = world.org("2nd Pine")
    pine = world.unit(other, "Pine cubs")
    admin = client_for(org["admin"].user, org["org"])
    assert admin.get(URL, {"fields": "full_name", "units": pine.pk}).status_code == 400


def test_text_is_never_a_formula(org, client_for):
    org["jad"].first_name = '=HYPERLINK("http://evil.example","click")'
    org["jad"].save()
    res = client_for(org["admin"].user, org["org"]).get(URL, {"fields": "first_name,date_of_birth"})
    ws = load_workbook(BytesIO(res.content)).active
    cell = next(c for c in ws["A"] if c.value and str(c.value).startswith("=HYPERLINK"))
    assert cell.data_type == "s"
