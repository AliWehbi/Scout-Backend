"""
CSV / Excel import of members.

The flow is stateless: the client uploads the file for a preview (validation + duplicate
detection) and uploads it again with the confirmed column mapping and per-row decisions
to commit. Nothing is written during preview.
"""

import csv
import io
import json
from datetime import date, datetime

from django.db import transaction
from django.db.models.functions import Lower
from django.utils.translation import gettext as _
from rest_framework.exceptions import ValidationError

from apps.units.models import Unit, UnitMembership, UnitPosition

from . import services
from .models import CustomField, Gender, Guardian, GuardianLink, MedicalInfo, Member, MemberKind, Relationship

MAX_ROWS = 2000

# Built-in import targets; custom fields are added as "custom.<key>".
TARGETS = [
    "first_name",
    "last_name",
    "name_ar",
    "date_of_birth",
    "gender",
    "kind",
    "unit",
    "position",
    "phone",
    "email",
    "address",
    "guardian_name",
    "guardian_phone",
    "guardian_email",
    "guardian_relationship",
    "allergies",
]

SYNONYMS = {
    "first_name": ["first name", "firstname", "first", "given name", "name", "الاسم", "الاسم الأول"],
    "last_name": ["last name", "lastname", "surname", "family name", "اسم العائلة", "الشهرة"],
    "name_ar": ["arabic name", "name (arabic)", "الاسم بالعربية"],
    "date_of_birth": ["dob", "date of birth", "birth date", "birthdate", "birthday", "تاريخ الميلاد"],
    "gender": ["gender", "sex", "الجنس"],
    "kind": ["type", "member type", "kind"],
    "unit": ["unit", "section", "group", "الوحدة", "المرحلة"],
    "position": ["position", "role in unit", "الموقع"],
    "phone": ["phone", "mobile", "telephone", "الهاتف"],
    "email": ["email", "e-mail", "البريد"],
    "address": ["address", "العنوان"],
    "guardian_name": ["parent", "guardian", "parent name", "guardian name", "ولي الأمر", "اسم ولي الأمر"],
    "guardian_phone": ["parent mobile", "parent phone", "guardian phone", "هاتف ولي الأمر"],
    "guardian_email": ["parent email", "guardian email", "بريد ولي الأمر"],
    "guardian_relationship": ["relationship", "relation", "صلة القرابة"],
    "allergies": ["allergies", "allergy", "الحساسية"],
}

DATE_FORMATS = ["%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d.%m.%Y", "%m/%d/%Y"]


def _norm(text) -> str:
    return " ".join(str(text or "").replace("_", " ").strip().lower().split())


def read_table(upload) -> tuple[list[str], list[list[str]]]:
    name = (upload.name or "").lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(upload, read_only=True, data_only=True).active
        rows = [
            [("" if c is None else (c.date().isoformat() if isinstance(c, datetime) else str(c))) for c in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    else:
        raw = upload.read()
        text = None
        for encoding in ("utf-8-sig", "cp1256", "latin-1"):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t") if text.strip() else csv.excel
        rows = list(csv.reader(io.StringIO(text), dialect))
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ValidationError({"file": [_("The file is empty.")]})
    if len(rows) - 1 > MAX_ROWS:
        raise ValidationError({"file": [_("Import at most %(n)s rows at a time.") % {"n": MAX_ROWS}]})
    headers = [str(h).strip() for h in rows[0]]
    body = [[str(c).strip() for c in r] + [""] * (len(headers) - len(r)) for r in rows[1:]]
    return headers, body


def suggest_mapping(headers: list[str], custom_fields) -> dict[str, str]:
    """header -> target ('' means not imported)."""
    lookup = {}
    for target, words in SYNONYMS.items():
        for word in words + [target]:
            lookup.setdefault(_norm(word), target)
    for cf in custom_fields:
        for word in (cf.key, cf.label_en, cf.label_ar):
            if word:
                lookup.setdefault(_norm(word), f"custom.{cf.key}")
    mapping, used = {}, set()
    for header in headers:
        target = lookup.get(_norm(header), "")
        if target and target in used:
            target = ""
        mapping[header] = target
        if target:
            used.add(target)
    return mapping


def _parse_date(value: str):
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


class Importer:
    def __init__(self, organization, upload, mapping: dict | None):
        self.org = organization
        self.headers, self.rows = read_table(upload)
        self.custom_fields = {cf.key: cf for cf in organization_custom_fields(organization)}
        self.mapping = mapping or suggest_mapping(self.headers, self.custom_fields.values())
        valid = set(TARGETS) | {f"custom.{k}" for k in self.custom_fields}
        bad = [t for t in self.mapping.values() if t and t not in valid]
        if bad:
            raise ValidationError({"mapping": [_("Unknown target: %(t)s") % {"t": ", ".join(bad)}]})
        self.units = {}
        for unit in Unit.objects.filter(organization=organization):
            for name in (unit.name_en, unit.name_ar, unit.section):
                if name:
                    self.units[_norm(name)] = unit
        self.existing = {}
        for m in Member.objects.filter(organization=organization).annotate(
            lf=Lower("first_name"), ll=Lower("last_name")
        ):
            self.existing.setdefault((m.lf.strip(), m.ll.strip()), []).append(m)

    def _values(self, raw: list[str]) -> dict:
        return {target: raw[i] for i, (header, target) in enumerate(self.mapping.items()) if target and i < len(raw)}

    def validate_row(self, values: dict) -> tuple[dict, dict]:
        """Returns (cleaned, errors) for one row."""
        cleaned, errors = {}, {}
        first = values.get("first_name", "").strip()
        if not first:
            errors["first_name"] = _("First name is required.")
        cleaned["first_name"] = first
        cleaned["last_name"] = values.get("last_name", "").strip()
        cleaned["name_ar"] = values.get("name_ar", "").strip()
        if values.get("date_of_birth"):
            dob = _parse_date(values["date_of_birth"])
            if dob is None or dob > date.today():
                errors["date_of_birth"] = _("Date isn’t valid — use DD/MM/YYYY.")
            cleaned["date_of_birth"] = dob
        gender = _norm(values.get("gender"))
        if gender:
            if gender in ("m", "male", "boy", "ذكر"):
                cleaned["gender"] = Gender.MALE
            elif gender in ("f", "female", "girl", "أنثى"):
                cleaned["gender"] = Gender.FEMALE
            else:
                errors["gender"] = _("Use male or female.")
        kind = _norm(values.get("kind"))
        cleaned["kind"] = MemberKind.ADULT if kind in ("adult", "leader", "قائد") else MemberKind.YOUTH
        unit_name = values.get("unit", "").strip()
        if unit_name:
            unit = self.units.get(_norm(unit_name))
            if unit is None:
                errors["unit"] = _("No unit called “%(name)s” — pick one or create it.") % {"name": unit_name}
            cleaned["unit"] = unit
        position = _norm(values.get("position")).replace(" ", "_")
        cleaned["position"] = position if position in UnitPosition.values else UnitPosition.MEMBER
        for key in ("phone", "email", "address", "guardian_name", "guardian_phone", "guardian_email"):
            cleaned[key] = values.get(key, "").strip()
        rel = _norm(values.get("guardian_relationship"))
        cleaned["guardian_relationship"] = rel if rel in Relationship.values else Relationship.GUARDIAN
        allergies = [a.strip() for a in values.get("allergies", "").replace(";", ",").split(",") if a.strip()]
        cleaned["allergies"] = [{"name": a, "severity": "moderate", "notes": ""} for a in allergies]
        custom = {}
        for target, value in values.items():
            if target.startswith("custom.") and value:
                field = self.custom_fields[target.split(".", 1)[1]]
                try:
                    custom[field.key] = services.clean_custom_value(field, value)
                except ValidationError as exc:
                    detail = exc.detail[0] if isinstance(exc.detail, list) else exc.detail
                    errors[target] = str(detail)
        cleaned["custom_data"] = custom
        return cleaned, errors

    def find_duplicate(self, cleaned: dict):
        key = (cleaned["first_name"].lower(), cleaned["last_name"].lower())
        for m in self.existing.get(key, []):
            if not cleaned.get("date_of_birth") or not m.date_of_birth or m.date_of_birth == cleaned["date_of_birth"]:
                return m
        return None

    def preview(self) -> dict:
        rows, counts = [], {"total": 0, "ready": 0, "errors": 0, "duplicates": 0}
        for index, raw in enumerate(self.rows, start=2):
            values = self._values(raw)
            cleaned, errors = self.validate_row(values)
            duplicate = None if errors else self.find_duplicate(cleaned)
            status = "error" if errors else "duplicate" if duplicate else "ready"
            counts["total"] += 1
            counts[{"error": "errors", "duplicate": "duplicates", "ready": "ready"}[status]] += 1
            rows.append(
                {
                    "row": index,
                    "status": status,
                    "values": values,
                    "errors": errors,
                    "duplicate_of": (
                        {
                            "id": duplicate.pk,
                            "full_name": duplicate.full_name,
                            "date_of_birth": duplicate.date_of_birth,
                        }
                        if duplicate
                        else None
                    ),
                }
            )
        samples = {
            header: next((row[i] for row in self.rows if i < len(row) and row[i]), "")
            for i, header in enumerate(self.headers)
        }
        return {
            "columns": self.headers,
            "samples": samples,
            "mapping": self.mapping,
            "targets": self.targets(),
            "summary": counts,
            "rows": rows,
        }

    def targets(self) -> list[dict]:
        return [{"key": t, "label": t} for t in TARGETS] + [
            {"key": f"custom.{cf.key}", "label": cf.label_en, "label_ar": cf.label_ar}
            for cf in self.custom_fields.values()
        ]

    @transaction.atomic
    def commit(self, decisions: dict, actor) -> dict:
        created = updated = skipped = 0
        guardians_by_phone = {g.phone: g for g in Guardian.objects.filter(organization=self.org).exclude(phone="")}
        for index, raw in enumerate(self.rows, start=2):
            cleaned, errors = self.validate_row(self._values(raw))
            if errors:
                skipped += 1
                continue
            duplicate = self.find_duplicate(cleaned)
            decision = decisions.get(str(index), "skip" if duplicate else "create")
            if decision == "skip":
                skipped += 1
                continue
            if duplicate and decision == "update":
                member = duplicate
                for field in ("last_name", "name_ar", "date_of_birth", "gender", "phone", "email", "address"):
                    if cleaned.get(field):
                        setattr(member, field, cleaned[field])
                member.custom_data = {**(member.custom_data or {}), **cleaned["custom_data"]}
                member.save()
                updated += 1
            else:
                member = Member.objects.create(
                    organization=self.org,
                    **{
                        k: cleaned[k] or ("" if k not in ("date_of_birth",) else None)
                        for k in (
                            "first_name",
                            "last_name",
                            "name_ar",
                            "date_of_birth",
                            "phone",
                            "email",
                            "address",
                        )
                        if k in cleaned
                    },
                    gender=cleaned.get("gender", ""),
                    kind=cleaned["kind"],
                    custom_data=cleaned["custom_data"],
                )
                created += 1
            if cleaned.get("unit"):
                UnitMembership.objects.update_or_create(
                    member=member, unit=cleaned["unit"], defaults={"position": cleaned["position"]}
                )
            if cleaned["guardian_name"]:
                guardian = guardians_by_phone.get(cleaned["guardian_phone"]) if cleaned["guardian_phone"] else None
                if guardian is None:
                    first, _sep, last = cleaned["guardian_name"].partition(" ")
                    guardian = Guardian.objects.create(
                        organization=self.org,
                        first_name=first,
                        last_name=last,
                        phone=cleaned["guardian_phone"],
                        email=cleaned["guardian_email"],
                    )
                    if guardian.phone:
                        guardians_by_phone[guardian.phone] = guardian
                GuardianLink.objects.get_or_create(
                    member=member,
                    guardian=guardian,
                    defaults={
                        "relationship": cleaned["guardian_relationship"],
                        "is_primary": not member.guardian_links.exists(),
                    },
                )
            if cleaned["allergies"]:
                medical, _created = MedicalInfo.objects.get_or_create(member=member)
                names = {a["name"].lower() for a in medical.allergies}
                medical.allergies = medical.allergies + [
                    a for a in cleaned["allergies"] if a["name"].lower() not in names
                ]
                medical.updated_by = actor
                medical.save()
        return {"created": created, "updated": updated, "skipped": skipped}


def organization_custom_fields(organization):
    return CustomField.objects.filter(organization=organization, is_active=True)


def parse_json_field(raw, name):
    if raw in (None, ""):
        return None
    if isinstance(raw, dict):
        return raw
    try:
        value = json.loads(raw)
    except (TypeError, ValueError):
        raise ValidationError({name: [_("Invalid JSON.")]}) from None
    if not isinstance(value, dict):
        raise ValidationError({name: [_("Expected an object.")]})
    return value
