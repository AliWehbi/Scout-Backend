import io
import json

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from apps.customfields.models import CustomField
from apps.events.models import Event, EventInvitation, EventType
from apps.events.services import create_default_event_types
from apps.members.models import Guardian, GuardianLink, Member
from apps.organizations.models import Organization, make_org_code
from apps.units.models import UnitMembership

FIELDS = "/api/v1/custom-fields/"
EVENTS = "/api/v1/events/"
TEMPLATES = "/api/v1/document-templates/"


@pytest.fixture(autouse=True)
def _private_root(settings, tmp_path):
    settings.PRIVATE_MEDIA_ROOT = tmp_path / "private"
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def setup(world, client_for):
    org = world.org("1st Cedar Scout Group")
    create_default_event_types(org)
    cubs = world.unit(org, "Cubs")
    admin = world.member(org, world.user("rania"), owner=True)
    karim = world.member(org, world.user("karim"), roles=[("unit_leader", cubs)])
    types = {t.name_en: t for t in EventType.objects.filter(organization=org)}
    return dict(org=org, cubs=cubs, admin=admin, karim=karim, types=types, rania=client_for(admin.user, org))


def add_field(client, **data):
    body = {"entity": "event", "label_en": "Meeting point", "field_type": "text", **data}
    return client.post(FIELDS, body, format="json")


def event_body(setup, **extra):
    return {
        "event_type": setup["types"]["Activity"].pk,
        "title": "Autumn camp",
        "start_date": "2030-10-23",
        "end_date": "2030-10-25",
        "units": [setup["cubs"].pk],
        **extra,
    }


def make_docx(*paragraphs) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def text_of(data: bytes) -> str:
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_organization_code(world):
    taken = set()

    def code(name):
        c = make_org_code(name, lambda x: x in taken)
        taken.add(c)
        return c

    assert code("1st Cedar Scout Group") == "cedar"
    assert code("Cedar Scouts") == "cedar2"
    assert code("Scouts of Lebanon") == "lebanon"
    assert code("فوج الأرز") == "org"
    assert code("2nd Group") == "org2nd"
    org = world.org("1st Cedar Scout Group")
    assert org.code.startswith("cedar")
    # It never changes, even when the organization is renamed.
    first = org.code
    org.name = "Something else"
    org.save()
    org.refresh_from_db()
    assert org.code == first
    assert Organization._meta.get_field("code").editable is False


def test_event_fields_codes_and_rules(setup, client_for):
    rania, org = setup["rania"], setup["org"]
    res = add_field(rania, code="meeting_point", default="Scout hall", event_types=[setup["types"]["Meeting"].pk])
    assert res.status_code == 201, res.data
    assert res.data["template_code"] == f"event.{org.code}_meeting_point"
    assert res.data["entity"] == "event" and res.data["visible_to_parents"] is False  # hidden from families by default
    assert res.data["event_types"] == [setup["types"]["Meeting"].pk]

    # Codes are optional, checked, and unique per kind of record.
    assert add_field(rania, label_en="Budget approved", field_type="yesno").data["template_code"] == ""
    assert add_field(rania, label_en="X", code="meeting_point").status_code == 400
    assert add_field(rania, label_en="X", code="Meeting Point").status_code == 400
    assert add_field(rania, label_en="X", code="9lives").status_code == 400
    # The default must fit the type.
    assert add_field(rania, label_en="Seats", field_type="number", default="many").status_code == 400
    assert add_field(rania, label_en="Bus", field_type="choice", options=["A", "B"], default="C").status_code == 400

    # Lists are per kind of record; members stay the default for older callers.
    assert [f["label_en"] for f in rania.get(FIELDS).data] == []
    assert len(rania.get(f"{FIELDS}?entity=event").data) == 2
    # A field can't switch kind.
    field_id = res.data["id"]
    rania.patch(f"{FIELDS}{field_id}/", {"entity": "member"}, format="json")
    assert CustomField.objects.get(pk=field_id).entity == "event"

    # Only "Manage organization settings" changes fields; anyone in the organization can read them.
    karim = client_for(setup["karim"].user, org)
    assert add_field(karim, label_en="Nope").status_code == 403
    assert karim.get(f"{FIELDS}?entity=event").status_code == 200


def test_event_values(setup):
    rania, types = setup["rania"], setup["types"]
    camp_only = add_field(
        rania, label_en="Bus company", code="bus", required=True, event_types=[types["Meeting"].pk]
    ).data
    everyone = add_field(rania, label_en="Contact email", code="contact_email", default="camp@cedar.org").data
    yes_no = add_field(rania, label_en="Budget approved", field_type="yesno").data

    # New events start with the defaults; a required field only matters for its event types.
    res = rania.post(EVENTS, event_body(setup), format="json")
    assert res.status_code == 201, res.data
    assert res.data["custom_data"] == {everyone["key"]: "camp@cedar.org"}
    event_id = res.data["id"]

    assert rania.post(EVENTS, event_body(setup, event_type=types["Meeting"].pk), format="json").status_code == 400
    camp = rania.post(
        EVENTS,
        event_body(setup, event_type=types["Meeting"].pk, custom_data={camp_only["key"]: "Nakhoul"}),
        format="json",
    )
    assert camp.status_code == 201 and camp.data["custom_data"][camp_only["key"]] == "Nakhoul"

    # Values are checked; unknown or non-applying fields are refused; edits keep the other values.
    url = f"{EVENTS}{event_id}/"
    assert rania.patch(url, {"custom_data": {yes_no["key"]: "maybe"}}, format="json").status_code == 400
    assert rania.patch(url, {"custom_data": {camp_only["key"]: "x"}}, format="json").status_code == 400
    assert rania.patch(url, {"custom_data": {"nope": "x"}}, format="json").status_code == 400
    res = rania.patch(url, {"custom_data": {yes_no["key"]: "yes"}}, format="json")
    assert res.data["custom_data"] == {everyone["key"]: "camp@cedar.org", yes_no["key"]: True}
    # Switching to a type with a required field needs its value.
    assert rania.patch(url, {"event_type": types["Meeting"].pk}, format="json").status_code == 400


def test_families_see_only_shown_fields(setup, client_for, world):
    rania, org = setup["rania"], setup["org"]
    shown = add_field(rania, label_en="Meeting point", default="Scout hall", visible_to_parents=True).data
    add_field(rania, label_en="Insurance no.", default="POL-1")
    add_field(rania, label_en="Pick-up", field_type="date", visible_to_parents=True)  # no value: left out
    event = Event.objects.create(
        organization=org, event_type=setup["types"]["Activity"], title="Hike", start_date="2030-10-23",
        end_date="2030-10-23", status="published",
    )  # fmt: skip
    jad = Member.objects.create(organization=org, first_name="Jad")
    UnitMembership.objects.create(member=jad, unit=setup["cubs"])
    maya = world.member(org, world.user("maya"), roles=[("parent", None)])
    guardian = Guardian.objects.create(organization=org, first_name="Maya", membership=maya)
    GuardianLink.objects.create(member=jad, guardian=guardian)
    EventInvitation.objects.create(event=event, member=jad, include_parents=True)
    data = client_for(maya.user, org).get("/api/v1/my/events/").data
    assert data[0]["custom_fields"] == [
        {"key": shown["key"], "label_en": "Meeting point", "label_ar": "Meeting point", "value_en": "Scout hall",
         "value_ar": "Scout hall"}
    ]  # fmt: skip


def test_documents_use_custom_codes(setup):
    rania, org = setup["rania"], setup["org"]
    point = add_field(rania, label_en="Meeting point", code="meeting_point", default="Scout hall").data
    add_field(rania, label_en="Bus company", code="bus")
    add_field(rania, label_en="Budget approved", field_type="yesno", code="budget_ok")
    prefix = f"event.{org.code}_"

    # The placeholder list starts with the organization's own codes.
    fields = rania.get(f"{TEMPLATES}fields/").data[0]["fields"]
    assert [f["name"] for f in fields if f.get("custom")] == [
        f"{prefix}meeting_point",
        f"{prefix}bus",
        f"{prefix}budget_ok",
    ]

    # Templates may use them; a code that doesn't exist is refused like any unknown field.
    bad = make_docx("{{ " + prefix + "typo }}")
    res = rania.post(TEMPLATES, {"kind": "event_sheet", "name": "Bad", "file": SimpleUploadedFile("b.docx", bad)})
    assert res.status_code == 400
    doc = make_docx("Meet at {{ " + prefix + "meeting_point }}", "Bus: {{ " + prefix + "bus }}")
    res = rania.post(
        TEMPLATES, {"kind": "event_sheet", "name": "Camp sheet", "file": SimpleUploadedFile("c.docx", doc)}
    )
    assert res.status_code == 201, res.data
    template = res.data["id"]

    event = rania.post(EVENTS, event_body(setup, custom_data={point["key"]: "Achrafieh hall"}), format="json").data
    url = f"{EVENTS}{event['id']}/"

    # Before exporting: the fields this template uses, with where each value comes from.
    rows = rania.get(f"{url}document-values/?kind=event_sheet&template={template}").data
    assert [(r["code"], r["value"], r["source"]) for r in rows] == [
        ("meeting_point", "Achrafieh hall", "record"),
        ("bus", None, "empty"),
    ]
    # The built-in templates don't use custom codes, so there's nothing to check.
    assert rania.get(f"{url}document-values/?kind=event_sheet&template=").data == []

    # Values edited for this document go into the file only; the event keeps its own.
    values = json.dumps({"bus": "Nakhoul Transport"})
    res = rania.get(f"{url}export/", {"kind": "event_sheet", "template": template, "lang": "en", "values": values})
    assert res.status_code == 200
    text = text_of(b"".join(res.streaming_content))
    assert "Meet at Achrafieh hall" in text and "Bus: Nakhoul Transport" in text
    assert "bus" not in Event.objects.get(pk=event["id"]).custom_data
    bad_values = json.dumps({"budget_ok": "perhaps"})
    res = rania.get(f"{url}export/", {"kind": "event_sheet", "template": template, "values": bad_values})
    assert res.status_code == 400


def test_member_fields_take_codes_and_defaults(setup):
    rania, org = setup["rania"], setup["org"]
    # Member documents are coming: member fields already take codes, in their own namespace.
    res = rania.post(
        FIELDS,
        {"label_en": "School", "code": "school", "default": "Cedar High", "applies_to": "youth"},
        format="json",
    )
    assert res.status_code == 201, res.data
    assert res.data["template_code"] == f"member.{org.code}_school"
    assert rania.post(FIELDS, {"label_en": "X", "code": "school"}, format="json").status_code == 400
    # The same code is fine on another kind of record.
    assert add_field(rania, label_en="School trip", code="school").status_code == 201
    # Member documents don't exist yet, so event templates don't list member codes.
    names = [f["name"] for f in rania.get(f"{TEMPLATES}fields/").data[0]["fields"] if f.get("custom")]
    assert names == [f"event.{org.code}_school"]

    # New members start with the defaults of the fields that apply to them.
    key = res.data["key"]
    youth = rania.post("/api/v1/members/", {"first_name": "Jad", "kind": "youth"}, format="json")
    assert youth.status_code == 201, youth.data
    assert Member.objects.get(pk=youth.data["id"]).custom_data == {key: "Cedar High"}
    adult = rania.post("/api/v1/members/", {"first_name": "Karim", "kind": "adult"}, format="json")
    assert Member.objects.get(pk=adult.data["id"]).custom_data == {}
    # Existing members keep what they have.
    member = Member.objects.get(pk=youth.data["id"])
    rania.patch(f"/api/v1/members/{member.pk}/", {"custom_data": {key: "Other school"}}, format="json")
    member.refresh_from_db()
    assert member.custom_data == {key: "Other school"}
