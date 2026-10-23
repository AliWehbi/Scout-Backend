import io
import os
import subprocess
import zipfile

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from docx import Document

from apps.audit.models import AuditLog
from apps.documents.models import DocumentTemplate
from apps.documents.rendering import RenderError, RenderJob, render_docx
from apps.events.models import Event, EventInvitation, EventType
from apps.events.services import create_default_event_types
from apps.members.models import Guardian, GuardianLink, Member
from apps.units.models import UnitMembership

URL = "/api/v1/document-templates/"


@pytest.fixture(autouse=True)
def _private_root(settings, tmp_path):
    settings.PRIVATE_MEDIA_ROOT = tmp_path / "private"
    settings.MEDIA_ROOT = tmp_path / "media"


@pytest.fixture
def setup(world):
    org = world.org()
    create_default_event_types(org)
    cubs, scouts = world.unit(org, "Cubs"), world.unit(org, "Scouts")
    admin = world.member(org, world.user("rania"), owner=True)
    karim = world.member(org, world.user("karim"), roles=[("unit_leader", cubs)])
    nour = world.member(org, world.user("nour"), roles=[("assistant_leader", cubs)])
    sami = world.member(org, world.user("sami"), roles=[("unit_leader", scouts)])
    event = Event.objects.create(
        organization=org,
        event_type=EventType.objects.get(organization=org, name_en="Activity"),
        title="Autumn camp",
        start_date="2030-10-23",
        end_date="2030-10-25",
        requires_consent=True,
        consent_text="I allow my child to go.",
        status="published",
    )
    event.units.set([cubs])
    jad = Member.objects.create(organization=org, first_name="Jad", first_name_ar="جاد")
    UnitMembership.objects.create(member=jad, unit=cubs)
    EventInvitation.objects.create(event=event, member=jad, response="going")
    guardian = Guardian.objects.create(organization=org, first_name="Maya", last_name="Saliba", phone="+961 3 111")
    GuardianLink.objects.create(member=jad, guardian=guardian, relationship="mother")
    return dict(org=org, admin=admin, karim=karim, nour=nour, sami=sami, event=event, jad=jad)


def make_docx(*paragraphs) -> bytes:
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def rezip(data: bytes, *, add=None, replace=None, compress=zipfile.ZIP_DEFLATED) -> bytes:
    """Copy a .docx, swapping or adding parts, to build malicious variants."""
    add, replace = add or {}, replace or {}
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(out, "w", compress) as dst:
        for info in src.infolist():
            content = src.read(info)
            if info.filename in replace:
                content = replace[info.filename](content)
            dst.writestr(info.filename, content)
        for name, content in add.items():
            dst.writestr(name, content)
    return out.getvalue()


def upload(client, data: bytes, *, kind="event_sheet", name="Our sheet", filename="sheet.docx", **extra):
    file = SimpleUploadedFile(filename, data, content_type="application/octet-stream")
    return client.post(URL, {"kind": kind, "name": name, "file": file, **extra}, format="multipart")


def text_of(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        parts += [c.text for row in table.rows for c in row.cells]
    return "\n".join(parts)


def file_error(res) -> str:
    return " ".join(res.data["errors"]["file"])


# ------------------------------------------------------------------ uploading


def test_upload_lands_in_private_storage(setup, client_for, settings):
    admin = client_for(setup["admin"].user, setup["org"])
    res = upload(admin, make_docx("{{ event.title }} – {{ event.type_ar }}"), filename="../../etc/passwd.docx")
    assert res.status_code == 201, res.data
    assert res.data["is_default"] is True  # first one of its kind
    assert res.data["original_filename"] == "passwd.docx"
    template = DocumentTemplate.objects.get(pk=res.data["id"])
    path = template.file.path
    assert path.startswith(str(settings.PRIVATE_MEDIA_ROOT))
    assert "passwd" not in path
    assert not (settings.MEDIA_ROOT / "templates").exists()
    assert AuditLog.objects.filter(action="document_template.uploaded").exists()

    res = admin.get(f"{URL}{template.pk}/download/")
    assert res.status_code == 200
    assert res["Content-Type"].startswith("application/vnd.openxmlformats")
    assert b"".join(res.streaming_content) == template.file.open("rb").read()


@pytest.mark.parametrize(
    "variant, message",
    [
        ("not_zip", "not a valid Word document"),
        ("macro_type", "macros"),
        ("vba", "macros"),
        ("ole", "embedded objects"),
        ("activex", "macros or controls"),
        ("remote_template", "linked templates"),
        ("external_image", "loads content from the internet"),
        ("doctype", "not a valid Word document"),
        ("zip_bomb", "too large once unpacked"),
        ("traversal", "not a valid Word document"),
    ],
)
def test_unsafe_uploads_are_rejected(setup, client_for, variant, message):
    base = make_docx("{{ event.title }}")
    rels = "word/_rels/document.xml.rels"

    def add_rel(rel_type, target, external=True):
        mode = ' TargetMode="External"' if external else ""

        def edit(xml: bytes) -> bytes:
            rel = (
                f'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                f'relationships/{rel_type}" Target="{target}"{mode}/>'
            )
            return xml.replace(b"</Relationships>", rel.encode() + b"</Relationships>")

        return edit

    data = {
        "not_zip": lambda: b"%PDF-1.7 definitely not a docx",
        "macro_type": lambda: rezip(
            base,
            replace={
                "[Content_Types].xml": lambda x: x.replace(b"document.main+xml", b"document.macroEnabled.main+xml")
            },
        ),
        "vba": lambda: rezip(base, add={"word/vbaProject.bin": b"\x00" * 64}),
        "ole": lambda: rezip(base, add={"word/embeddings/oleObject1.bin": b"\x00" * 64}),
        "activex": lambda: rezip(base, add={"word/activeX/activeX1.xml": b"<ax/>"}),
        "remote_template": lambda: rezip(
            base, replace={rels: add_rel("attachedTemplate", "http://evil.example/t.dotm")}
        ),
        "external_image": lambda: rezip(base, replace={rels: add_rel("image", "http://evil.example/pixel.png")}),
        "doctype": lambda: rezip(
            base,
            replace={
                "word/document.xml": lambda x: x.replace(
                    b"?>", b'?><!DOCTYPE d [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>', 1
                )
            },
        ),
        "zip_bomb": lambda: rezip(base, add={"word/media/bomb.xml": b"0" * (60 * 1024 * 1024)}),
        "traversal": lambda: rezip(base, add={"../../evil.xml": b"<x/>"}),
    }[variant]()

    res = upload(client_for(setup["admin"].user, setup["org"]), data)
    assert res.status_code == 400, res.data
    assert message in file_error(res)
    assert not DocumentTemplate.objects.exists()


def test_plain_hyperlinks_are_allowed(setup, client_for):
    base = make_docx("{{ event.title }}")
    rel = (
        b'<Relationship Id="rIdX" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/'
        b'hyperlink" Target="https://scouts.example" TargetMode="External"/></Relationships>'
    )
    data = rezip(base, replace={"word/_rels/document.xml.rels": lambda x: x.replace(b"</Relationships>", rel)})
    assert upload(client_for(setup["admin"].user, setup["org"]), data).status_code == 201


def test_upload_size_and_extension(setup, client_for):
    admin = client_for(setup["admin"].user, setup["org"])
    res = upload(admin, make_docx("x"), filename="sheet.docm")
    assert res.status_code == 400 and "(.docx)" in file_error(res)
    res = upload(admin, b"0" * (5 * 1024 * 1024 + 1))
    assert res.status_code == 400 and "5 MB" in file_error(res)


@pytest.mark.parametrize(
    "text, message",
    [
        ("{% for p in participants %}", "mistake in a placeholder"),
        ("{{ secret_key }}", "Unknown field: secret_key"),
        ("{{ event.titel }}", "Unknown field"),
        ("{{ rows }}", "Unknown field: rows"),  # belongs to the attendance sheet
        ("{{ ''.__class__.__mro__ }}", "isn't allowed"),
        ("{{ event.title.__init__.__globals__ }}", "isn't allowed"),
        ("{{ 'a' * 100000000 }}", "isn't allowed"),
        ("{{ 9 ** 999999 }}", "isn't allowed"),
        ("{% for i in range(1000000) %}x{% endfor %}", "isn't allowed"),
        ("{{ 'a'.ljust(1000000000) }}", "isn't allowed"),
    ],
)
def test_template_mistakes_and_sandbox_escapes_are_rejected(setup, client_for, text, message):
    res = upload(client_for(setup["admin"].user, setup["org"]), make_docx(text))
    assert res.status_code == 400, res.data
    assert message in file_error(res)


def test_runaway_template_is_stopped():
    loops = "{% for a in range(1000) %}{% for b in range(1000) %}{% for c in range(1000) %}{% endfor %}{% endfor %}{% endfor %}"
    with pytest.raises(RenderError) as exc:
        render_docx(RenderJob(make_docx(loops), {}))
    assert exc.value.code in ("too_large", "timeout")


def test_values_are_xml_escaped(setup, client_for):
    jad = setup["jad"]
    jad.first_name = '</w:t></w:r><w:r><w:t>INJECTED & "quoted"'
    jad.save()
    admin = client_for(setup["admin"].user, setup["org"])
    res = admin.get(f"/api/v1/events/{setup['event'].pk}/export/", {"kind": "event_sheet", "lang": "en"})
    assert res.status_code == 200
    data = b"".join(res.streaming_content)
    assert jad.first_name in text_of(data)  # shown as text, not turned into markup
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        assert b"<w:t>INJECTED" not in zf.read("word/document.xml")


def test_render_process_gets_no_secrets(monkeypatch):
    monkeypatch.setenv("DJANGO_SECRET_KEY", "top-secret")
    monkeypatch.setenv("DATABASE_URL", "postgres://u:p@db/x")
    seen = {}
    real_run = subprocess.run

    def spy(*args, **kwargs):
        seen.update(kwargs["env"])
        return real_run(*args, **kwargs)

    monkeypatch.setattr(subprocess, "run", spy)
    render_docx(RenderJob(make_docx("hi"), {}))
    assert "DJANGO_SECRET_KEY" not in seen and "DATABASE_URL" not in seen


# ------------------------------------------------------------------ managing templates


def test_default_switches_and_delete_removes_file(setup, client_for, django_capture_on_commit_callbacks):
    admin = client_for(setup["admin"].user, setup["org"])
    first = upload(admin, make_docx("{{ event.title }}"), name="A").data
    second = upload(admin, make_docx("{{ event.title }}"), name="B", is_default=True).data
    assert second["is_default"] is True
    assert DocumentTemplate.objects.get(pk=first["id"]).is_default is False

    res = admin.patch(f"{URL}{first['id']}/", {"is_default": True, "name": "A2"}, format="json")
    assert res.status_code == 200 and res.data["name"] == "A2"
    assert DocumentTemplate.objects.get(pk=second["id"]).is_default is False

    template = DocumentTemplate.objects.get(pk=second["id"])
    path = template.file.path
    with django_capture_on_commit_callbacks(execute=True):
        assert admin.delete(f"{URL}{template.pk}/").status_code == 204
    assert not os.path.exists(path)


def test_fields_and_starters(setup, client_for):
    admin = client_for(setup["admin"].user, setup["org"])
    kinds = {k["kind"]: k for k in admin.get(f"{URL}fields/").data}
    assert set(kinds) == {"event_sheet", "attendance_sheet", "consent_slips", "members_list"}
    names = [f["name"] for f in kinds["consent_slips"]["fields"]]
    assert "event.title" in names and "slips" in names
    for kind in kinds:
        for lang in ("en", "ar"):
            res = admin.get(f"{URL}starter/", {"kind": kind, "lang": lang})
            assert res.status_code == 200
            # every built-in template must pass the same checks as an upload
            data = b"".join(res.streaming_content)
            ext = "xlsx" if kind == "members_list" else "docx"
            res = upload(admin, data, kind=kind, name=f"{kind} {lang}", filename=f"starter.{ext}")
            assert res.status_code == 201, res.data
    assert admin.get(f"{URL}starter/", {"kind": "secrets"}).status_code == 400


def test_permissions_and_tenancy(setup, world, client_for):
    karim = client_for(setup["karim"].user, setup["org"])
    assert karim.get(URL).status_code == 200  # can pick a template when exporting
    assert upload(karim, make_docx("{{ event.title }}")).status_code == 403
    assert karim.get(f"{URL}fields/").status_code == 403

    admin = client_for(setup["admin"].user, setup["org"])
    mine = upload(admin, make_docx("{{ event.title }}")).data

    other = world.org("2nd Pine")
    stranger = world.member(other, world.user("lina"), owner=True)
    lina = client_for(stranger.user, other)
    assert lina.get(URL).data == []
    assert lina.get(f"{URL}{mine['id']}/download/").status_code == 404
    assert lina.delete(f"{URL}{mine['id']}/").status_code == 404
    # nor can it be used to export one of their own events
    create_default_event_types(other)
    theirs = Event.objects.create(
        organization=other,
        event_type=EventType.objects.get(organization=other, name_en="Meeting"),
        title="Pine meeting",
        start_date="2030-10-23",
        end_date="2030-10-23",
    )
    res = lina.get(f"/api/v1/events/{theirs.pk}/export/", {"kind": "event_sheet", "template": mine["id"]})
    assert res.status_code == 404


# ------------------------------------------------------------------ exporting


def export(client, event, **params):
    return client.get(f"/api/v1/events/{event.pk}/export/", params)


def test_export_each_kind_with_built_in_templates(setup, client_for):
    admin = client_for(setup["admin"].user, setup["org"])
    event = setup["event"]

    res = export(admin, event, kind="event_sheet", lang="en")
    assert res.status_code == 200
    assert "attachment" in res["Content-Disposition"] and "Autumn camp - Event sheet.docx" in res["Content-Disposition"]
    assert res["Cache-Control"] == "no-store"
    text = text_of(b"".join(res.streaming_content))
    assert "Autumn camp" in text and "Jad" in text and "Going" in text and "1st Cedar" in text

    res = export(admin, event, kind="event_sheet", lang="ar")
    text = text_of(b"".join(res.streaming_content))
    assert "جاد" in text and "سيحضر" in text and "أكتوبر" in text
    assert "filename*=utf-8''" in res["Content-Disposition"]  # Arabic name, RFC 5987

    res = export(admin, event, kind="attendance_sheet", date="2030-10-24")
    assert res.status_code == 200 and "2030-10-24" in res["Content-Disposition"]
    assert "24 October 2030" in text_of(b"".join(res.streaming_content))
    assert export(admin, event, kind="attendance_sheet", date="2030-11-01").status_code == 400

    res = export(admin, event, kind="consent_slips")
    text = text_of(b"".join(res.streaming_content))
    assert "Maya Saliba" in text and "I allow my child to go." in text
    assert AuditLog.objects.filter(action="event.exported").count() == 4


def test_export_uses_uploaded_default_or_chosen_template(setup, client_for):
    admin = client_for(setup["admin"].user, setup["org"])
    event = setup["event"]
    custom = upload(
        admin,
        make_docx(
            "CUSTOM {{ event.title }} / {{ event.type_ar }}",
            "{%p for p in participants %}",
            "{{ p.n }}. {{ p.name_ar }} {{ p.signed_text }}",
            "{%p endfor %}",
        ),
    ).data
    text = text_of(b"".join(export(admin, event, kind="event_sheet").streaming_content))
    assert "CUSTOM Autumn camp / نشاط" in text and "1. جاد No" in text

    text = text_of(b"".join(export(admin, event, kind="event_sheet", template="").streaming_content))
    assert "CUSTOM" not in text  # empty template = built-in

    text = text_of(b"".join(export(admin, event, kind="event_sheet", template=custom["id"]).streaming_content))
    assert "CUSTOM" in text
    # a template of another kind can't be used
    assert export(admin, event, kind="consent_slips", template=custom["id"]).status_code == 404


def test_tampered_template_file_is_refused(setup, client_for):
    admin = client_for(setup["admin"].user, setup["org"])
    custom = upload(admin, make_docx("{{ event.title }}")).data
    template = DocumentTemplate.objects.get(pk=custom["id"])
    with open(template.file.path, "wb") as fh:
        fh.write(make_docx("{{ event.title }} changed"))
    res = export(admin, setup["event"], kind="event_sheet")
    assert res.status_code == 400 and "changed on the server" in res.data["detail"]


def test_export_permissions(setup, client_for):
    event = setup["event"]
    karim = client_for(setup["karim"].user, setup["org"])  # manages Cubs
    for kind in ("event_sheet", "attendance_sheet", "consent_slips"):
        assert export(karim, event, kind=kind).status_code == 200

    nour = client_for(setup["nour"].user, setup["org"])  # assistant: view + check-in
    assert export(nour, event, kind="event_sheet").status_code == 200
    assert export(nour, event, kind="attendance_sheet").status_code == 200
    assert export(nour, event, kind="consent_slips").status_code == 403  # guardians' phones

    sami = client_for(setup["sami"].user, setup["org"])  # Scouts leader, can't see a Cubs event
    assert export(sami, event, kind="event_sheet").status_code == 404

    assert export(karim, event, kind="nope").status_code == 400
