"""
Excel templates for lists (the members export first).

An organization designs a sheet in Excel: logo, organization details, titles, colours, a footer.
Codes in cells are replaced as plain text (no formulas or logic run): `{{ org.name }}`,
`{{ org.cedar_phone }}`, `{{ generated_at }}`... `{{ org.logo }}` alone in a cell becomes the
logo, and `{{ table }}` alone in a cell marks where the data goes: the header row is written
there (in that cell's style) and the rows below it (in the style of the cell under it). Anything
the template has below the marker moves down to make room.
"""

import re
from copy import copy
from io import BytesIO

from django.utils.translation import gettext as _
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XlImage
from openpyxl.utils import get_column_letter
from rest_framework.exceptions import ValidationError

CODE_RE = re.compile(r"\{\{\s*([A-Za-z_][A-Za-z0-9_.]*)\s*\}\}")
TABLE = "table"
LOGO = "org.logo"
LOGO_HEIGHT = 64  # pixels
FORMULA_START = ("=", "+", "-", "@")


def _reject(message):
    raise ValidationError({"file": [message]})


def _load(data: bytes):
    try:
        return load_workbook(BytesIO(data), rich_text=False)
    except Exception:  # noqa: BLE001 - any parsing problem means it's not a usable workbook
        _reject(_("This file is not a valid Excel workbook."))


def _codes(wb):
    """(sheet, cell, [codes]) for every text cell that contains codes."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and "{{" in cell.value:
                    names = CODE_RE.findall(cell.value)
                    if names:
                        yield ws, cell, names


def used_codes(data: bytes) -> set[str]:
    return {name for _ws, _cell, names in _codes(_load(data)) for name in names}


def check_template(data: bytes, allowed: set[str]) -> None:
    """Unknown codes, and a missing or misplaced {{ table }}, are refused at upload."""
    wb = _load(data)
    tables = []
    unknown = set()
    for _ws, cell, names in _codes(wb):
        for name in names:
            if name not in allowed:
                unknown.add(name)
            if name in (TABLE, LOGO) and not CODE_RE.fullmatch(cell.value.strip()):
                _reject(_("Put {{ %(code)s }} alone in its cell.") % {"code": name})
        if TABLE in names:
            tables.append(cell)
    if unknown:
        _reject(_("Unknown fields: %(list)s") % {"list": ", ".join(sorted(unknown))})
    if not tables:
        _reject(_("Put {{ table }} in the cell where the list should start."))
    if len(tables) > 1:
        _reject(_("Use {{ table }} only once."))


def _style(cell) -> dict:
    return {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
    }


def apply_style(cell, style: dict | None) -> None:
    if not style:
        return
    for name, value in style.items():
        setattr(cell, name, copy(value) if name != "number_format" else value)


def set_text(cell, text: str) -> None:
    """Text that stays text: a value starting with "=" must never become a formula."""
    cell.value = text
    cell.data_type = "s"
    if text[:1] in FORMULA_START:
        cell.quotePrefix = True


def _shift_below(ws, row: int, amount: int) -> None:
    """Make room under the marker row: cells, merged areas, row heights and pictures move down."""
    if amount <= 0:
        return
    below_merged = [m for m in ws.merged_cells.ranges if m.min_row > row]
    for merged in below_merged:
        ws.merged_cells.remove(merged)
    heights = {r: ws.row_dimensions[r].height for r in range(row + 1, ws.max_row + 1) if ws.row_dimensions[r].height}
    ws.insert_rows(row + 1, amount)
    for merged in below_merged:
        merged.shift(row_shift=amount)
        ws.merged_cells.add(merged)
    for r in heights:
        ws.row_dimensions[r].height = None
    for r, height in heights.items():
        ws.row_dimensions[r + amount].height = height
    for image in ws._images:
        anchor = getattr(image, "anchor", None)
        start = getattr(anchor, "_from", None)
        if start is not None and start.row >= row:  # 0-based: rows below the marker row
            start.row += amount
            end = getattr(anchor, "to", None)
            if end is not None:
                end.row += amount


def _logo(cell, ws, png: bytes | None) -> None:
    cell.value = None
    if not png:
        return
    image = XlImage(BytesIO(png))
    if image.height:
        ratio = LOGO_HEIGHT / image.height
        image.height, image.width = LOGO_HEIGHT, max(int(image.width * ratio), 1)
    ws.add_image(image, cell.coordinate)


def render(data: bytes, values: dict[str, str], logo: bytes | None, write_table) -> bytes:
    """
    Fill a checked template. `write_table(ws, row, col, header_style, body_style, make_room)` writes
    the header in row `row` and the data below it, calling `make_room(n)` first so that whatever the
    template has under the marker moves down by the `n` data rows.
    """
    wb = _load(data)
    marker = None
    for ws, cell, names in list(_codes(wb)):
        if TABLE in names:
            marker = (ws, cell)
            continue
        if LOGO in names:
            _logo(cell, ws, logo)
            continue
        set_text(cell, CODE_RE.sub(lambda m: values.get(m.group(1), ""), cell.value))
    if marker is None:
        _reject(_("Put {{ table }} in the cell where the list should start."))
    ws, cell = marker
    row, col = cell.row, cell.column
    header_style = _style(cell)
    below = ws.cell(row=row + 1, column=col)
    body_style = _style(below) if below.has_style else None
    cell.value = None
    write_table(ws, row, col, header_style, body_style, lambda n: _shift_below(ws, row, n))
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def column_widths(ws, col: int, widths: list[int]) -> None:
    """Size the table's columns to their content, unless the template set a width."""
    for offset, width in enumerate(widths):
        dim = ws.column_dimensions[get_column_letter(col + offset)]
        if not dim.customWidth:
            dim.width = min(max(width + 2, 8), 50)


STARTER = {
    "en": {"title": "Members list", "exported": "Exported on", "by": "by", "count": "Members:", "sheet": "Members"},
    "ar": {
        "title": "قائمة الأعضاء",
        "exported": "تاريخ التصدير",
        "by": "بواسطة",
        "count": "عدد الأعضاء:",
        "sheet": "الأعضاء",
    },
}


def starter(lang: str) -> bytes:
    """A starting point to design in Excel: logo, organization, title, export details, the table."""
    from openpyxl.styles import Alignment, Font, PatternFill

    t = STARTER["ar" if lang == "ar" else "en"]
    wb = Workbook()
    ws = wb.active
    ws.title = t["sheet"]
    ws.sheet_view.rightToLeft = lang == "ar"
    ws.row_dimensions[1].height = 52
    ws["A1"] = "{{ org.logo }}"
    ws["B1"] = "{{ org.name }}"
    ws["B1"].font = Font(bold=True, size=16)
    ws["B1"].alignment = Alignment(vertical="center")
    ws["B2"] = t["title"]
    ws["B2"].font = Font(size=13, color="595D53")
    ws["B3"] = f"{t['exported']} {{{{ generated_at }}}} {t['by']} {{{{ generated_by }}}}"
    ws["B3"].font = Font(size=10, color="85887E")
    ws["B4"] = f"{t['count']} {{{{ count }}}}"
    ws["B4"].font = Font(size=10, color="85887E")
    ws.column_dimensions["A"].width = 14
    ws["A6"] = "{{ table }}"
    ws["A6"].font = Font(bold=True, color="FFFFFF")
    ws["A6"].fill = PatternFill("solid", fgColor="2F5D3A")
    ws["A7"].alignment = Alignment(vertical="top")
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
