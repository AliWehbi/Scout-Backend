# ruff: noqa: UP031 - %-formatting reads better than str.format around Jinja braces
"""
Builds the built-in Word templates (one per kind and language) next to this file.

    python -m apps.documents.defaults.build_defaults

The generated .docx files are committed; rerun this after changing it. They double as the
"starter" templates admins download, edit in Word and upload again.
"""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

HERE = Path(__file__).resolve().parent
FONT = "Arial"
MUTED = RGBColor(0x6B, 0x72, 0x80)

T = {
    "en": {
        "type": "Type",
        "dates": "Dates",
        "location": "Location",
        "units": "Units",
        "leaders": "Leaders",
        "mandatory": "Mandatory",
        "cost": "Cost",
        "reply_by": "Reply by",
        "summary": "Invited {{ counts.invited }} · Going {{ counts.going }} · Maybe {{ counts.maybe }} · "
        "Not going {{ counts.not_going }} · No reply {{ counts.no_reply }}",
        "participants": "Participants",
        "name": "Name",
        "unit": "Unit",
        "reply": "Reply",
        "consent": "Consent",
        "attendance": "Attendance",
        "patrol": "Patrol",
        "status": "Status",
        "signature": "Signature / note",
        "att_summary": "Present {{ counts.present }} · Absent {{ counts.absent }} · Excused {{ counts.excused }}",
        "consent_title": "Parent consent",
        "member": "Member",
        "guardians": "Parent / guardian",
        "signed": "Signed by {{ s.signed_by }} on {{ s.signed_at }}",
        "sign_here": "Signature: ____________________     Date: ______________",
        "footer": "Exported {{ generated_at }} by {{ generated_by }}",
        "type_field": "event.type_en",
        "units_field": "event.unit_names_en",
        "leaders_field": "event.leader_names",
        "name_field": "full_name",
    },
    "ar": {
        "type": "النوع",
        "dates": "التاريخ",
        "location": "المكان",
        "units": "الوحدات",
        "leaders": "القادة",
        "mandatory": "إلزامي",
        "cost": "التكلفة",
        "reply_by": "آخر موعد للرد",
        "summary": "المدعوون {{ counts.invited }} · سيحضر {{ counts.going }} · ربما {{ counts.maybe }} · "
        "لن يحضر {{ counts.not_going }} · لم يرد {{ counts.no_reply }}",
        "participants": "المشاركون",
        "name": "الاسم",
        "unit": "الوحدة",
        "reply": "الرد",
        "consent": "الموافقة",
        "attendance": "الحضور",
        "patrol": "الطليعة",
        "status": "الحالة",
        "signature": "التوقيع / ملاحظة",
        "att_summary": "حاضر {{ counts.present }} · غائب {{ counts.absent }} · معذور {{ counts.excused }}",
        "consent_title": "موافقة ولي الأمر",
        "member": "العضو",
        "guardians": "ولي الأمر",
        "signed": "وقّعها {{ s.signed_by }} بتاريخ {{ s.signed_at }}",
        "sign_here": "التوقيع: ____________________     التاريخ: ______________",
        "footer": "صدر في {{ generated_at }} بواسطة {{ generated_by }}",
        "type_field": "event.type_ar",
        "units_field": "event.unit_names_ar",
        "leaders_field": "event.leader_names_ar",
        "name_field": "name_ar",
    },
}


# ------------------------------------------------------------------ helpers


def _set(el, tag, **attrs):
    child = el.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        el.append(child)
    for key, value in attrs.items():
        child.set(qn(key), value)
    return child


def new_document(rtl: bool) -> Document:
    doc = Document()
    for style in doc.styles:
        if style.type != 1:  # paragraph styles only
            continue
        rpr = style.element.get_or_add_rPr()
        _set(rpr, "w:rFonts", **{"w:ascii": FONT, "w:hAnsi": FONT, "w:cs": FONT})
        if rtl:
            _set(rpr, "w:rtl")
            _set(style.element.get_or_add_pPr(), "w:bidi")
    doc.styles["Normal"].font.size = Pt(10.5)
    section = doc.sections[0]
    if rtl:
        _set(section._sectPr, "w:bidi")
    return doc


def styled_run(p, text, *, bold=False, size=None, muted=False):
    run = p.add_run(text)
    rpr = run._r.get_or_add_rPr()
    if bold:
        run.bold = True
        _set(rpr, "w:bCs")  # Arabic is "complex script" text with its own bold and size
    if size:
        run.font.size = Pt(size)
        _set(rpr, "w:szCs", **{"w:val": str(size * 2)})
    if muted:
        run.font.color.rgb = MUTED
    return run


def para(doc, text="", *, bold=False, size=None, muted=False):
    p = doc.add_paragraph()
    if text:
        styled_run(p, text, bold=bold, size=size, muted=muted)
    return p


def table(doc, rows: int, cols: int, rtl: bool):
    t = doc.add_table(rows=rows, cols=cols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    if rtl:
        _set(t._tbl.tblPr, "w:bidiVisual")
    return t


def cell(t, r, c, text, *, bold=False):
    cell_ = t.cell(r, c)
    cell_.text = ""
    styled_run(cell_.paragraphs[0], text, bold=bold)
    return cell_


def header(doc, lang, title):
    para(doc, "{{ org.logo }}")
    para(doc, "{{ org.name }}", bold=True, muted=True)
    para(doc, title, bold=True, size=18)


def footer(doc, lang):
    styled_run(doc.sections[0].footer.paragraphs[0], T[lang]["footer"], size=8, muted=True)


def loop_table(doc, rtl, headings, loop_open, values, loop_close="{%tr endfor %}"):
    """A table whose middle row repeats: heading row, `{%tr for %}` row, item row, `{%tr endfor %}` row."""
    t = table(doc, 4, len(headings), rtl)
    for c, text in enumerate(headings):
        cell(t, 0, c, text, bold=True)
    cell(t, 1, 0, loop_open)
    for c, text in enumerate(values):
        cell(t, 2, c, text)
    cell(t, 3, 0, loop_close)
    return t


# ------------------------------------------------------------------ documents


def details(doc, lang, rtl):
    s = T[lang]
    rows = [
        (s["type"], "{{ %s }}" % s["type_field"]),
        (
            s["dates"],
            "{{ event.start_date }} {{ event.start_time }}"
            "{% if event.end_date != event.start_date or event.end_time %} – {{ event.end_date }} {{ event.end_time }}{% endif %}",
        ),
        (s["location"], "{{ event.location }}"),
        (s["units"], "{{ %s }}" % s["units_field"]),
        (s["leaders"], "{{ %s }}" % s["leaders_field"]),
        (s["mandatory"], "{{ event.is_mandatory_text }}"),
        (s["cost"], "{% if event.cost %}{{ event.cost }} {{ event.currency }}{% endif %}"),
        (s["reply_by"], "{{ event.reply_deadline }}"),
    ]
    t = table(doc, len(rows), 2, rtl)
    for r, (key, value) in enumerate(rows):
        cell(t, r, 0, key, bold=True)
        cell(t, r, 1, value)


def event_sheet(lang):
    rtl, s = lang == "ar", T[lang]
    doc = new_document(rtl)
    header(doc, lang, "{{ event.title }}")
    details(doc, lang, rtl)
    para(doc)
    para(doc, "{{ event.description }}")
    para(doc, s["participants"], bold=True, size=13)
    para(doc, s["summary"], muted=True)
    name = s["name_field"]
    loop_table(
        doc,
        rtl,
        ["#", s["name"], s["unit"], s["reply"], s["consent"]],
        "{%tr for p in participants %}",
        ["{{ p.n }}", "{{ p.%s }}" % name, "{{ p.unit }}", "{{ p.response }}", "{{ p.signed_text }}"],
    )
    footer(doc, lang)
    return doc


def attendance_sheet(lang):
    rtl, s = lang == "ar", T[lang]
    doc = new_document(rtl)
    header(doc, lang, "%s – {{ event.title }}" % s["attendance"])
    para(doc, "{{ day }} · {{ event.location }}", muted=True)
    para(doc, "%s: {{ %s }}" % (s["leaders"], s["leaders_field"]))
    para(doc, s["att_summary"], muted=True)
    name = s["name_field"]
    loop_table(
        doc,
        rtl,
        ["#", s["name"], s["patrol"], s["reply"], s["status"], s["signature"]],
        "{%tr for r in rows %}",
        ["{{ r.n }}", "{{ r.%s }}" % name, "{{ r.patrol }}", "{{ r.response }}", "{{ r.status }}", ""],
    )
    footer(doc, lang)
    return doc


def consent_slips(lang):
    rtl, s = lang == "ar", T[lang]
    doc = new_document(rtl)
    name = s["name_field"]
    para(doc, "{%p for s in slips %}")
    header(doc, lang, "%s – {{ event.title }}" % s["consent_title"])
    para(doc, "{{ event.start_date }} – {{ event.end_date }} · {{ event.location }}", muted=True)
    t = table(doc, 2, 2, rtl)
    cell(t, 0, 0, s["member"], bold=True)
    cell(t, 0, 1, "{{ s.%s }}" % name)
    cell(t, 1, 0, s["guardians"], bold=True)
    cell(t, 1, 1, "{{ s.guardian_names }}")
    para(doc)
    para(doc, "{{ s.consent_text }}")
    para(doc)
    para(doc, "{% if s.signed %}" + s["signed"] + "{% else %}" + s["sign_here"] + "{% endif %}")
    para(doc, "{%p if not loop.last %}")
    doc.add_page_break()
    para(doc, "{%p endif %}")
    para(doc, "{%p endfor %}")
    footer(doc, lang)
    return doc


BUILDERS = {"event_sheet": event_sheet, "attendance_sheet": attendance_sheet, "consent_slips": consent_slips}


def main():
    for kind, build in BUILDERS.items():
        for lang in ("en", "ar"):
            doc = build(lang)
            path = HERE / f"{kind}.{lang}.docx"
            doc.save(path)
            print(path.name)


if __name__ == "__main__":
    main()
